import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import numpy as np
import re
import csv
import hashlib
from io import BytesIO
from pathlib import Path
from datetime import datetime, time
from fpdf import FPDF
import smtplib
from email.message import EmailMessage
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter


class ReportPDF(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", size=8)
        self.set_text_color(120, 120, 120)
        self.cell(
            0,
            6,
            "Developed by : Dhruv Pathak and Rahul Singh",
            0,
            0,
            "R"
        )


# =========================================================
# PAGE CONFIG
# =========================================================
st.set_page_config(
    page_title="Meteorological Operations Suite",
    layout="wide",
    initial_sidebar_state="collapsed"
)

st.markdown("""
<style>
.dev-credit {
    position: fixed;
    right: 18px;
    bottom: 14px;
    z-index: 9999;
    background: rgba(255, 255, 255, 0.88);
    color: #4B5563;
    padding: 6px 10px;
    border-radius: 10px;
    font-size: 12px;
    border: 1px solid #D9E2EC;
    box-shadow: 0 2px 8px rgba(0,0,0,0.08);
    backdrop-filter: blur(6px);
}
</style>

<div class="dev-credit">
    Developed by : Dhruv Pathak and Rahul Singh
</div>
""", unsafe_allow_html=True)

# =========================================================
# STYLING
# =========================================================
st.markdown("""
<style>
.block-container {
    padding-top: 2.4rem;
    padding-bottom: 1rem;
    max-width: 97%;
}
.hero-panel {
    margin-top: 0.8rem;
}
:root {
    --bg: #EFF4F9;
    --surface: #FFFFFF;
    --surface-soft: #F8FBFE;
    --line: #D9E2EC;
    --line-soft: #E9EEF5;
    --text: #172B4D;
    --muted: #6B778C;
    --title: #0F1C2E;
    --primary: #3FA9F5;
    --primary-dark: #1F5E8C;
    --shadow-sm: 0 1px 2px rgba(15, 23, 42, 0.05);
    --shadow-md: 0 10px 24px rgba(15, 23, 42, 0.07);
    --radius-lg: 22px;
}
html, body, [class*="css"] {
    font-family: "Segoe UI", Arial, sans-serif;
    color: var(--text);
}
.stApp {
    background:
        linear-gradient(180deg, rgba(63,169,245,0.12) 0%, rgba(239,244,249,1) 240px),
        var(--bg);
}
#MainMenu, footer {
    visibility: visible !important;
}
.main-shell {
    margin-bottom: 0.8rem;
}
.hero-panel {
    background: linear-gradient(135deg, rgba(255,255,255,0.98) 0%, rgba(249,252,255,0.99) 100%);
    border: 1px solid rgba(217,226,236,0.95);
    border-radius: 28px;
    padding: 20px 22px 16px 22px;
    box-shadow: var(--shadow-md);
    margin-bottom: 0.75rem;
}
.hero-title {
    font-size: 1.5rem;
    font-weight: 760;
    color: var(--title);
    margin-bottom: 0.15rem;
    letter-spacing: -0.02em;
}
.hero-subtitle {
    font-size: 0.93rem;
    color: var(--muted);
    line-height: 1.5;
    max-width: 980px;
    margin-bottom: 0.75rem;
}
.pill {
    display: inline-block;
    padding: 6px 12px;
    border-radius: 999px;
    font-size: 0.75rem;
    font-weight: 650;
    margin-right: 8px;
    margin-top: 4px;
    border: 1px solid var(--line-soft);
    background: #FFFFFF;
    color: var(--text);
}
.section-label {
    font-size: 1rem;
    font-weight: 720;
    color: var(--title);
    margin: 0.2rem 0 0.45rem 0;
}
.subtle-note {
    background: rgba(255,255,255,0.78);
    border: 1px solid var(--line);
    border-left: 4px solid var(--primary);
    color: var(--text);
    padding: 10px 12px;
    border-radius: 12px;
    font-size: 0.88rem;
    margin-bottom: 0.8rem;
    box-shadow: var(--shadow-sm);
}
.site-shell {
    background: linear-gradient(180deg, rgba(255,255,255,0.94) 0%, rgba(252,253,254,0.98) 100%);
    border: 1px solid var(--line);
    border-radius: 22px;
    padding: 12px 12px 10px 12px;
    box-shadow: var(--shadow-md);
    margin-bottom: 0.9rem;
}
.site-title {
    font-size: 1.04rem;
    font-weight: 740;
    color: var(--title);
    margin-bottom: 0.08rem;
}
.site-subtitle {
    color: var(--muted);
    font-size: 0.82rem;
}
.metric-chip {
    display: inline-block;
    padding: 4px 10px;
    border-radius: 999px;
    font-size: 0.73rem;
    font-weight: 650;
    margin-right: 6px;
    margin-top: 6px;
    border: 1px solid var(--line-soft);
    background: #FFFFFF;
    color: #334155;
}
.range-badge {
    display: inline-block;
    padding: 4px 10px;
    border-radius: 999px;
    background: rgba(63,169,245,0.10);
    color: var(--primary-dark);
    font-size: 0.74rem;
    font-weight: 700;
    border: 1px solid rgba(63,169,245,0.18);
}
.helper-note {
    background: #F8FBFE;
    border: 1px dashed #C8D9EA;
    color: #46627F;
    padding: 9px 11px;
    border-radius: 12px;
    font-size: 0.82rem;
    margin: 0.25rem 0 0.7rem 0;
}
.small-panel {
    background: linear-gradient(180deg, #FFFFFF 0%, #FCFDFE 100%);
    border: 1px solid var(--line);
    border-radius: 18px;
    padding: 10px 12px;
    box-shadow: var(--shadow-sm);
}
.small-title {
    font-size: 0.84rem;
    font-weight: 700;
    color: var(--title);
    margin-bottom: 0.35rem;
}
.small-caption {
    color: var(--muted);
    font-size: 0.76rem;
    margin-bottom: 0.45rem;
}
div[data-testid="stMetric"] {
    background: linear-gradient(180deg, #FFFFFF 0%, #FBFCFE 100%);
    border: 1px solid var(--line);
    border-radius: 16px;
    padding: 8px 10px;
    box-shadow: var(--shadow-sm);
}
div[data-testid="stMetricLabel"] {
    color: var(--muted) !important;
    font-weight: 650 !important;
}
div[data-testid="stMetricValue"] {
    color: var(--title) !important;
    font-weight: 760 !important;
}
div[data-testid="stDataFrame"],
div[data-testid="stExpander"] {
    border: 1px solid var(--line);
    border-radius: 14px;
    overflow: hidden;
    background: #FFFFFF;
}
.stButton > button,
.stDownloadButton > button,
.stFormSubmitButton > button {
    border-radius: 12px !important;
    border: 1px solid var(--line) !important;
    background: #FFFFFF !important;
    color: var(--text) !important;
    font-weight: 650 !important;
    box-shadow: none !important;
}
.stFormSubmitButton > button {
    background: linear-gradient(135deg, #52B6F6 0%, #2F95DA 100%) !important;
    color: white !important;
    border-color: #52B6F6 !important;
}
div[data-testid="stRadio"] > div {
    flex-direction: row;
    gap: 8px;
}
div[data-testid="stRadio"] label {
    background: #FFFFFF;
    border: 1px solid var(--line);
    padding: 6px 10px;
    border-radius: 10px;
}
hr {
    border: none;
    border-top: 1px solid var(--line-soft);
    margin: 0.55rem 0;
}
</style>
""", unsafe_allow_html=True)

# =========================================================
# PATHS
# =========================================================
BASE_DIR = Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
UPLOAD_DIR = BASE_DIR / "data_uploads"
PROCESSED_DIR = BASE_DIR / "processed"
REPORT_DIR = BASE_DIR / "Meteorological Reports" / "Daily"
LOG_FILE = BASE_DIR / "upload_log.csv"

for folder in [UPLOAD_DIR, PROCESSED_DIR, REPORT_DIR]:
    folder.mkdir(parents=True, exist_ok=True)

# =========================================================
# MASTER TEMPLATE CONFIG
# =========================================================
MASTER_TEMPLATE_NAME = "MASTER-file.xlsx"
MASTER_SHEET_NAME = "clean data"

# =========================================================
# CONSTANTS
# =========================================================
METRIC_PATTERN = re.compile(
    r"^(.*?)\s(Wind Speed|Wind Direction|Irradiance|Precipitation)\s\[(.*?)\]$",
    flags=re.IGNORECASE
)

COLOR_MAP = {
    "Wind Speed": "#58AEE8",
    "Wind Direction": "#F1A55B",
    "Irradiance": "#4CC9B0",
    "Precipitation": "#9B7BF2"
}

# =========================================================
# HELPERS
# =========================================================


def apply_clean_data_table(ws):
    if ws.max_row < 2 or ws.max_column < 1:
        return

    existing_table_names = list(ws.tables.keys())
    for table_name in existing_table_names:
        if table_name == "TblCleanData":
            del ws.tables[table_name]

    table_ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    tab = Table(displayName="TblCleanData", ref=table_ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    ws.add_table(tab)
master_template_path = Path("Master-file.xlsx")
wb = load_workbook(master_template_path)
clean_ws = wb["Clean Data"]

# clear old rows / write headers / write new data here

apply_clean_data_table(clean_ws)

wb.save(output_path)

def send_reports_email_gmail(to_emails, subject, body, reports):
    sender = st.secrets["EMAIL_SENDER"]
    password = st.secrets["EMAIL_PASSWORD"]

    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = ", ".join(to_emails)
    msg.set_content(body)

    for report in reports:
        pdf_path = report["path"]
        file_name = report["file_name"]

        with open(pdf_path, "rb") as f:
            pdf_data = f.read()

        msg.add_attachment(
            pdf_data,
            maintype="application",
            subtype="pdf",
            filename=file_name
        )

    with smtplib.SMTP("smtp.gmail.com", 587) as server:
        server.ehlo()
        server.starttls()
        server.ehlo()
        server.login(sender, password)
        server.send_message(msg)


def safe_file_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(name))


def get_file_hash(file_bytes: bytes) -> str:
    return hashlib.md5(file_bytes).hexdigest()


def append_upload_log(record: dict):
    df_new = pd.DataFrame([record])
    if LOG_FILE.exists():
        try:
            df_old = pd.read_csv(LOG_FILE)
            df_all = pd.concat([df_old, df_new], ignore_index=True)
        except Exception:
            df_all = df_new
    else:
        df_all = df_new
    df_all.to_csv(LOG_FILE, index=False)


def load_upload_log():
    if LOG_FILE.exists():
        return pd.read_csv(LOG_FILE)
    return pd.DataFrame()


def split_variable(col_name):
    text = str(col_name).strip()
    match = METRIC_PATTERN.match(text)
    if match:
        site = match.group(1).strip()
        metric = match.group(2).strip().title()
        unit = match.group(3).strip()
        return site, metric, unit
    return text, "Value", ""


def parse_semicolon_table(lines):
    clean_lines = [ln for ln in lines if str(ln).strip()]
    if not clean_lines:
        raise ValueError("No readable lines found in the uploaded file.")

    reader = csv.reader(clean_lines, delimiter=";")
    rows = [list(r) for r in reader if r]

    if len(rows) < 2:
        raise ValueError("The file does not contain enough rows to parse.")

    header = [str(h).strip() for h in rows[0]]
    parsed_rows = []

    for row in rows[1:]:
        row = [str(x).strip() for x in row]
        if len(row) < len(header):
            row = row + [None] * (len(header) - len(row))
        elif len(row) > len(header):
            row = row[:len(header)]
        parsed_rows.append(row)

    df = pd.DataFrame(parsed_rows, columns=header)

    begin_col = next((c for c in df.columns if str(c).lower().startswith("begin")), None)
    end_col = next((c for c in df.columns if str(c).lower().startswith("end")), None)

    if begin_col is None:
        raise ValueError("Could not find the 'begin' timestamp column.")

    df[begin_col] = pd.to_datetime(df[begin_col], dayfirst=True, errors="coerce")
    if end_col:
        df[end_col] = pd.to_datetime(df[end_col], dayfirst=True, errors="coerce")

    for col in df.columns:
        if col not in [begin_col, end_col]:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df = df.dropna(subset=[begin_col]).sort_values(begin_col).reset_index(drop=True)
    return df, begin_col, end_col


def wide_to_long(df, begin_col, end_col):
    id_vars = [begin_col] + ([end_col] if end_col else [])
    value_vars = [c for c in df.columns if c not in id_vars]
    long_df = df.melt(id_vars=id_vars, value_vars=value_vars, var_name="parameter", value_name="value")
    meta = long_df["parameter"].apply(lambda x: pd.Series(split_variable(x), index=["site", "metric", "unit"]))
    long_df = pd.concat([long_df, meta], axis=1)

    rename_map = {begin_col: "begin"}
    if end_col:
        rename_map[end_col] = "end"

    long_df = long_df.rename(columns=rename_map)
    long_df = long_df.dropna(subset=["begin", "value"]).sort_values(["site", "metric", "begin"]).reset_index(drop=True)
    return long_df


@st.cache_data(show_spinner=False)
def parse_uploaded_content(file_bytes: bytes, file_name: str):
    ext = Path(file_name).suffix.lower()

    if ext in [".xlsx", ".xls"]:
        xdf = pd.read_excel(BytesIO(file_bytes), header=None, dtype=str)
        lines = []
        for _, row in xdf.iterrows():
            vals = [str(v).strip() for v in row.tolist() if pd.notna(v) and str(v).strip() != ""]
            if not vals:
                continue
            if len(vals) == 1:
                cell_text = vals[0].replace("\r\n", "\n").replace("\r", "\n")
                lines.extend([ln.strip() for ln in cell_text.split("\n") if ln.strip()])
            else:
                lines.append(";".join(vals))
    else:
        text = file_bytes.decode("utf-8-sig", errors="ignore")
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

    wide_df, begin_col, end_col = parse_semicolon_table(lines)
    long_df = wide_to_long(wide_df, begin_col, end_col)
    return wide_df, long_df, begin_col, end_col


@st.cache_data(show_spinner=False)
def prepare_long_df(long_df):
    df = long_df.copy()
    df["site"] = df["site"].astype(str).str.strip()
    df["metric"] = df["metric"].astype(str).str.strip()
    df["unit"] = df["unit"].astype(str).replace("nan", "")
    return df.sort_values(["site", "metric", "begin"]).reset_index(drop=True)


def save_upload_bundle(uploaded_file, file_bytes, wide_df, long_df):
    ts = datetime.now()
    stamp = ts.strftime("%Y%m%d_%H%M%S")
    day_folder = UPLOAD_DIR / ts.strftime("%Y-%m-%d")
    day_folder.mkdir(parents=True, exist_ok=True)

    cleaned_name = safe_file_name(uploaded_file.name)
    base_name = Path(cleaned_name).stem

    raw_path = day_folder / f"{stamp}_{cleaned_name}"
    wide_path = PROCESSED_DIR / f"{stamp}_{base_name}_wide.csv"
    long_path = PROCESSED_DIR / f"{stamp}_{base_name}_long.csv"

    with open(raw_path, "wb") as f:
        f.write(file_bytes)

    wide_df.to_csv(wide_path, index=False)
    long_df.to_csv(long_path, index=False)

    append_upload_log({
        "upload_timestamp": ts.strftime("%Y-%m-%d %H:%M:%S"),
        "original_file_name": uploaded_file.name,
        "saved_raw_file": str(raw_path),
        "saved_wide_csv": str(wide_path),
        "saved_long_csv": str(long_path),
        "rows_wide": len(wide_df),
        "rows_long": len(long_df),
        "sites": long_df["site"].nunique(),
        "metrics": long_df["metric"].nunique(),
        "start_time": str(long_df["begin"].min()),
        "end_time": str(long_df["begin"].max()),
        "file_hash": get_file_hash(file_bytes)
    })


def normalize_series(s):
    s = pd.to_numeric(s, errors="coerce")
    s_min = s.min()
    s_max = s.max()
    if pd.isna(s_min) or pd.isna(s_max):
        return pd.Series(np.nan, index=s.index)
    if s_max == s_min:
        return pd.Series(50.0, index=s.index)
    return ((s - s_min) / (s_max - s_min)) * 100.0


def get_filtered_site_metric_df(site_df, selected_metrics):
    return (
        site_df[site_df["metric"].isin(selected_metrics)]
        .copy()
        .dropna(subset=["begin"])
        .sort_values("begin")
    )


def get_last_n_days_df(site_df, selected_metrics, n_days):
    df = get_filtered_site_metric_df(site_df, selected_metrics)
    if df.empty:
        return df
    max_ts = df["begin"].max()
    start_ts = max_ts - pd.Timedelta(days=n_days)
    return df[df["begin"] >= start_ts].copy()


def get_custom_range_df(site_df, selected_metrics, start_dt, end_dt):
    df = get_filtered_site_metric_df(site_df, selected_metrics)
    if df.empty or start_dt is None or end_dt is None:
        return df.copy()
    return df[(df["begin"] >= start_dt) & (df["begin"] <= end_dt)].copy()


def build_latest_table(df):
    if df.empty:
        return pd.DataFrame(columns=["metric", "unit", "latest_time", "latest_value"])

    latest_df = (
        df.sort_values("begin")
        .groupby(["metric", "unit"], as_index=False)
        .tail(1)[["metric", "unit", "begin", "value"]]
        .rename(columns={"begin": "latest_time", "value": "latest_value"})
        .sort_values("metric")
        .reset_index(drop=True)
    )
    latest_df["latest_value"] = latest_df["latest_value"].round(2)
    return latest_df


def build_window_summary(df):
    if df.empty:
        return pd.DataFrame(columns=["metric", "unit", "avg", "min", "max", "latest_time"])

    summary_df = (
        df.groupby(["metric", "unit"], as_index=False)
        .agg(avg=("value", "mean"), min=("value", "min"), max=("value", "max"), latest_time=("begin", "max"))
        .sort_values("metric")
        .reset_index(drop=True)
    )

    for col in ["avg", "min", "max"]:
        summary_df[col] = summary_df[col].round(2)
    return summary_df


def installation_status_from_row(row):
    if row["wind_avg"] <= 8.0 and row["wind_max"] <= 10.0 and row["rain_total"] <= 1.0:
        return "Good Window"
    if row["wind_avg"] <= 10.0 and row["wind_max"] <= 12.0 and row["rain_total"] <= 5.0:
        return "Caution Window"
    return "Avoid Window"


def installation_status_color(status):
    if status == "Good Window":
        return (46, 204, 113)
    if status == "Caution Window":
        return (241, 196, 15)
    return (231, 76, 60)


def installation_status_text_color(status):
    if status == "Caution Window":
        return (70, 60, 0)
    return (255, 255, 255)


def overall_installation_status(score):
    if score >= 70:
        return "Good Window"
    if score >= 45:
        return "Caution Window"
    return "Avoid Window"


def build_installation_daily_table(site_view_df):
    if site_view_df.empty:
        return pd.DataFrame()

    df = site_view_df.copy()
    df["metric"] = df["metric"].astype(str).str.strip()
    df["date"] = pd.to_datetime(df["begin"]).dt.date

    wind_df = df[df["metric"].str.lower() == "wind speed"].copy()
    if wind_df.empty:
        return pd.DataFrame()

    rain_df = df[df["metric"].str.lower() == "precipitation"].copy()

    wind_unit_mode = wind_df["unit"].dropna().astype(str).mode()
    wind_unit = wind_unit_mode.iloc[0] if len(wind_unit_mode) > 0 else ""

    rain_unit_mode = rain_df["unit"].dropna().astype(str).mode()
    rain_unit = rain_unit_mode.iloc[0] if len(rain_unit_mode) > 0 else ""

    daily = (
        wind_df.groupby("date", as_index=False)
        .agg(
            wind_avg=("value", "mean"),
            wind_max=("value", "max"),
            wind_min=("value", "min"),
            wind_std=("value", "std"),
            obs=("value", "count")
        )
    )

    if not rain_df.empty:
        rain_daily = rain_df.groupby("date", as_index=False).agg(rain_total=("value", "sum"))
        daily = daily.merge(rain_daily, on="date", how="left")
    else:
        daily["rain_total"] = 0.0

    daily["wind_std"] = daily["wind_std"].fillna(0.0)
    daily["rain_total"] = daily["rain_total"].fillna(0.0)

    daily["readiness_score"] = (
        100
        - np.clip((daily["wind_avg"] - 6.0) * 10.0, 0, 35)
        - np.clip((daily["wind_max"] - 8.0) * 8.0, 0, 35)
        - np.clip(daily["wind_std"] * 6.0, 0, 15)
        - np.clip(daily["rain_total"] * 5.0, 0, 20)
    ).clip(0, 100).round().astype(int)

    daily["installation_window"] = daily.apply(installation_status_from_row, axis=1)
    daily["wind_unit"] = wind_unit
    daily["rain_unit"] = rain_unit

    return daily.sort_values("date").reset_index(drop=True)


def add_installation_insight_block(pdf, site_view_df):
    daily = build_installation_daily_table(site_view_df)

    if daily.empty:
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(110, 110, 110)
        pdf.multi_cell(
            0,
            4.5,
            "Installation readiness insight unavailable because Wind Speed data is not present for this chart window."
        )
        pdf.set_text_color(0, 0, 0)
        return

    wind_unit = daily["wind_unit"].iloc[0] if "wind_unit" in daily.columns else ""
    rain_unit = daily["rain_unit"].iloc[0] if "rain_unit" in daily.columns and str(daily["rain_unit"].iloc[0]).strip() else "mm"

    best = daily.sort_values(
        ["readiness_score", "wind_avg", "rain_total"],
        ascending=[False, True, True]
    ).iloc[0]

    worst = daily.sort_values(
        ["readiness_score", "wind_avg", "rain_total"],
        ascending=[True, False, False]
    ).iloc[0]

    good_days = int((daily["installation_window"] == "Good Window").sum())
    caution_days = int((daily["installation_window"] == "Caution Window").sum())
    avoid_days = int((daily["installation_window"] == "Avoid Window").sum())

    overall_score = int(round(daily["readiness_score"].mean()))
    overall_status = overall_installation_status(overall_score)
    overall_color = installation_status_color(overall_status)

    if pdf.get_y() > 190:
        pdf.add_page()

    box_x = pdf.l_margin
    box_w = pdf.w - pdf.l_margin - pdf.r_margin
    start_y = pdf.get_y()
    box_h = 94

    pdf.set_draw_color(217, 226, 236)
    pdf.set_fill_color(248, 251, 254)
    pdf.rect(box_x, start_y, box_w, box_h, "DF")

    pdf.set_xy(box_x + 3, start_y + 3)
    pdf.set_font("Helvetica", "B", 11)
    pdf.set_text_color(23, 43, 77)
    pdf.cell(0, 6, "Erection & Installation Readiness", ln=True)

    pdf.set_x(box_x + 3)
    pdf.set_font("Helvetica", "", 7.5)
    pdf.set_text_color(100, 110, 120)
    pdf.multi_cell(
        box_w - 6,
        4,
        "Indicative weather-to-action view for erection planning. Validate against OEM crane limits, lift plan and site stop-work rules."
    )

    card_y = pdf.get_y() + 1
    card_w = 42
    card_h = 18

    pdf.set_fill_color(*overall_color)
    pdf.rect(box_x + 3, card_y, card_w, card_h, "F")
    pdf.set_draw_color(255, 255, 255)
    pdf.rect(box_x + 3, card_y, card_w, card_h)

    txt_r, txt_g, txt_b = installation_status_text_color(overall_status)
    pdf.set_text_color(txt_r, txt_g, txt_b)
    pdf.set_xy(box_x + 5, card_y + 3)
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(card_w - 4, 6, f"{overall_score}/100", ln=True, align="C")
    pdf.set_x(box_x + 5)
    pdf.set_font("Helvetica", "B", 8)
    pdf.cell(card_w - 4, 4, overall_status, ln=True, align="C")

    info_x = box_x + 50
    info_w = box_w - 53

    pdf.set_text_color(23, 43, 77)
    pdf.set_xy(info_x, card_y)
    pdf.set_font("Helvetica", "B", 8.5)
    pdf.cell(
        info_w,
        5,
        f"Best day: {pd.to_datetime(best['date']).strftime('%d %b %Y')}  |  Score {int(best['readiness_score'])}",
        ln=True
    )
    pdf.set_x(info_x)
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(
        info_w,
        4.5,
        f"Avg wind {best['wind_avg']:.1f} {wind_unit} | Peak wind {best['wind_max']:.1f} {wind_unit} | Rain {best['rain_total']:.1f} {rain_unit}",
        ln=True
    )

    pdf.set_x(info_x)
    pdf.set_font("Helvetica", "B", 8.5)
    pdf.cell(
        info_w,
        5,
        f"Worst day: {pd.to_datetime(worst['date']).strftime('%d %b %Y')}  |  Score {int(worst['readiness_score'])}",
        ln=True
    )
    pdf.set_x(info_x)
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(
        info_w,
        4.5,
        f"Avg wind {worst['wind_avg']:.1f} {wind_unit} | Peak wind {worst['wind_max']:.1f} {wind_unit} | Rain {worst['rain_total']:.1f} {rain_unit}",
        ln=True
    )

    pdf.set_x(info_x)
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(
        info_w,
        4.5,
        f"Window count: Good {good_days} | Caution {caution_days} | Avoid {avoid_days}",
        ln=True
    )

    strip_y = card_y + 24
    strip_x = box_x + 3
    strip_w = box_w - 6
    n = max(len(daily), 1)
    gap = 1
    tile_w = (strip_w - gap * (n - 1)) / n

    pdf.set_text_color(23, 43, 77)
    pdf.set_xy(strip_x, strip_y - 5)
    pdf.set_font("Helvetica", "B", 8)
    pdf.cell(strip_w, 4, "Daily readiness strip", ln=True)

    for i, row in daily.reset_index(drop=True).iterrows():
        fill = installation_status_color(row["installation_window"])
        x = strip_x + i * (tile_w + gap)

        pdf.set_fill_color(*fill)
        pdf.set_draw_color(255, 255, 255)
        pdf.rect(x, strip_y, tile_w, 8, "FD")

        if tile_w >= 10:
            rr, gg, bb = installation_status_text_color(row["installation_window"])
            pdf.set_text_color(rr, gg, bb)
            pdf.set_font("Helvetica", "B", 6)
            pdf.set_xy(x, strip_y + 1.4)
            pdf.cell(tile_w, 3, str(int(row["readiness_score"])), align="C")

    pdf.set_text_color(90, 100, 110)
    pdf.set_font("Helvetica", "", 6.5)
    label_step = max(1, int(np.ceil(len(daily) / 6)))

    for i, row in daily.reset_index(drop=True).iterrows():
        if i % label_step == 0 or i == len(daily) - 1:
            x = strip_x + i * (tile_w + gap)
            pdf.set_xy(x - 1, strip_y + 9)
            pdf.cell(tile_w + 2, 3.5, pd.to_datetime(row["date"]).strftime("%d %b"), align="C")

    legend_y = strip_y + 15
    legend_items = [
        ("Good Window", (46, 204, 113)),
        ("Caution Window", (241, 196, 15)),
        ("Avoid Window", (231, 76, 60)),
    ]

    legend_x = strip_x
    for label, color in legend_items:
        pdf.set_fill_color(*color)
        pdf.rect(legend_x, legend_y, 4, 4, "F")
        pdf.set_text_color(70, 80, 90)
        pdf.set_xy(legend_x + 6, legend_y - 0.8)
        pdf.set_font("Helvetica", "", 7)
        pdf.cell(28, 5, label)
        legend_x += 42

    note_y = legend_y + 7
    pdf.set_xy(strip_x, note_y)
    pdf.set_text_color(23, 43, 77)
    pdf.set_font("Helvetica", "", 8)

    note_lines = [
        "Green: prioritize heavy lifts and blade installation sequencing. Amber: use for standby, pre-assembly and internal logistics.",
        "Red: treat as no-lift / contingency unless site-approved limits and actual conditions support work."
    ]

    for line in note_lines:
        pdf.set_x(strip_x)
        pdf.multi_cell(box_w - 6, 3.2, f"- {line}")

    pdf.set_text_color(0, 0, 0)
    pdf.set_y(start_y + box_h + 2)


def build_combined_chart(df, site_name, scale_mode, chart_id=None, title_suffix=""):
    if df.empty:
        return None

    plot_df = df.copy().sort_values("begin")

    if scale_mode == "Normalized (0-100)":
        plot_df["plot_value"] = plot_df.groupby("metric")["value"].transform(normalize_series)
        y_title = "Normalized scale (0-100)"
        y_range = [0, 100]
    else:
        plot_df["plot_value"] = plot_df["value"]
        y_title = "Actual values"
        y_range = None

    fig = go.Figure()

    for metric in sorted(plot_df["metric"].dropna().unique()):
        mdf = plot_df[plot_df["metric"] == metric].copy()
        if mdf.empty:
            continue

        unit_mode = mdf["unit"].dropna().astype(str).mode()
        unit = unit_mode.iloc[0] if len(unit_mode) > 0 else ""
        label = f"{metric} [{unit}]" if str(unit).strip() else metric
        color = COLOR_MAP.get(metric, "#64748B")

        fig.add_trace(go.Scatter(
            x=mdf["begin"],
            y=mdf["plot_value"],
            mode="lines",
            name=label,
            line=dict(width=2.4, color=color, shape="spline", smoothing=0.35),
            connectgaps=True,
            hovertemplate=f"<b>{label}</b><br>Time: %{{x}}<br>Value: %{{customdata:.2f}}<extra></extra>",
            customdata=mdf["value"]
        ))

    full_title = f"{site_name} | {title_suffix}" if title_suffix else site_name

    fig.update_layout(
    title=full_title,
    template="plotly_white",
    height=520,
    margin=dict(l=10, r=10, t=35, b=110),
    legend=dict(
        orientation="h",
        yanchor="top",
        y=-0.22,
        xanchor="center",
        x=0.5,
        bgcolor="rgba(255,255,255,0.0)",
        font=dict(size=10)
    ),
    legend_title="",
    hovermode="x unified",
    paper_bgcolor="rgba(0,0,0,0)",
    plot_bgcolor="#FFFFFF",
    uirevision=chart_id or f"{site_name}_{scale_mode}",
    font=dict(color="#172B4D")
)

    fig.update_xaxes(
        title="Timestamp",
        showgrid=True,
        gridcolor="#E7EEF5",
        zeroline=False,
        showline=True,
        linecolor="#D9E2EC"
    )
    fig.update_yaxes(
        title=y_title,
        showgrid=True,
        gridcolor="#E7EEF5",
        zeroline=False,
        showline=True,
        linecolor="#D9E2EC"
    )

    if y_range is not None:
        fig.update_yaxes(range=y_range)

    return fig


def make_metric_summary(df_metric):
    if df_metric.empty:
        return pd.DataFrame()

    summary = (
        df_metric.groupby("site", as_index=False)
        .agg(
            count=("value", "count"),
            avg=("value", "mean"),
            min=("value", "min"),
            max=("value", "max"),
            std=("value", "std"),
            total=("value", "sum")
        )
    )

    latest = (
        df_metric.sort_values("begin")
        .groupby("site", as_index=False)
        .tail(1)[["site", "begin", "value"]]
        .rename(columns={"begin": "latest_time", "value": "latest_value"})
    )

    summary = summary.merge(latest, on="site", how="left")
    for col in ["avg", "min", "max", "std", "total", "latest_value"]:
        summary[col] = summary[col].round(2)

    return summary.sort_values("avg", ascending=False).reset_index(drop=True)


def build_insights(df_metric, metric_name, unit):
    insights = []
    if df_metric.empty:
        return insights

    avg_by_site = df_metric.groupby("site")["value"].mean().sort_values(ascending=False)
    max_by_site = df_metric.groupby("site")["value"].max().sort_values(ascending=False)
    std_by_site = df_metric.groupby("site")["value"].std().sort_values(ascending=False)

    if not avg_by_site.empty:
        insights.append(f"Highest average {metric_name.lower()}: {avg_by_site.index[0]} ({avg_by_site.iloc[0]:.2f} {unit})")
    if not max_by_site.empty:
        insights.append(f"Highest observed {metric_name.lower()}: {max_by_site.index[0]} ({max_by_site.iloc[0]:.2f} {unit})")

    std_non_na = std_by_site.dropna()
    if not std_non_na.empty:
        insights.append(f"Most variable {metric_name.lower()}: {std_non_na.index[0]} (std dev {std_non_na.iloc[0]:.2f} {unit})")

    if metric_name.lower() == "precipitation":
        total_by_site = df_metric.groupby("site")["value"].sum().sort_values(ascending=False)
        if not total_by_site.empty:
            insights.append(f"Highest cumulative precipitation: {total_by_site.index[0]} ({total_by_site.iloc[0]:.2f} {unit})")

    return insights


def add_chart_page(pdf, page_title, image_path, source_file_name, selected_metrics, scale_mode):
    pdf.add_page()
    pdf.set_font("Helvetica", style="B", size=13)
    pdf.cell(0, 8, page_title, ln=True)
    pdf.set_font("Helvetica", size=9)
    pdf.cell(0, 5, f"Source file: {source_file_name}", ln=True)
    pdf.cell(0, 5, f"Metrics: {', '.join(selected_metrics)}", ln=True)
    pdf.cell(0, 5, f"Scale mode: {scale_mode}", ln=True)
    pdf.cell(0, 5, f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
    pdf.ln(3)

    image_y = pdf.get_y()
    image_w = 190
    image_h = 107
    pdf.image(str(image_path), x=10, y=image_y, w=image_w)
    return image_y + image_h


def add_no_data_page(pdf, page_title, source_file_name, selected_metrics, scale_mode):
    pdf.add_page()
    pdf.set_font("Helvetica", style="B", size=13)
    pdf.cell(0, 8, page_title, ln=True)
    pdf.set_font("Helvetica", size=9)
    pdf.cell(0, 5, f"Source file: {source_file_name}", ln=True)
    pdf.cell(0, 5, f"Metrics: {', '.join(selected_metrics)}", ln=True)
    pdf.cell(0, 5, f"Scale mode: {scale_mode}", ln=True)
    pdf.ln(8)
    pdf.multi_cell(0, 6, "No data available for this chart window.")


def build_site_pdf_report_bytes(site_name, site_df, selected_metrics, scale_mode, source_file_name, report_timestamp):
    safe_site = safe_file_name(site_name)
    temp_dir = REPORT_DIR / "_temp_images"
    temp_dir.mkdir(parents=True, exist_ok=True)
    report_timestamp = datetime.now().strftime("%Y%m%d%H%M%S")

    chart_15_df = get_last_n_days_df(site_df, selected_metrics, 15)
    chart_2_df = get_last_n_days_df(site_df, selected_metrics, 2)

    fig_15 = build_combined_chart(
        df=chart_15_df,
        site_name=site_name,
        scale_mode=scale_mode,
        chart_id=f"{safe_site}_pdf_15_{report_timestamp}",
        title_suffix="Last 15 Days"
    )
    fig_2 = build_combined_chart(
        df=chart_2_df,
        site_name=site_name,
        scale_mode=scale_mode,
        chart_id=f"{safe_site}_pdf_2_{report_timestamp}",
        title_suffix="Last 2 Days"
    )

    img_15_path = temp_dir / f"{safe_site}_{report_timestamp}_15day.png"
    img_2_path = temp_dir / f"{safe_site}_{report_timestamp}_2day.png"
    created_images = []

    try:
        if fig_15 is not None:
            fig_15.write_image(str(img_15_path), format="png", width=1600, height=900)
            created_images.append(img_15_path)

        if fig_2 is not None:
            fig_2.write_image(str(img_2_path), format="png", width=1600, height=900)
            created_images.append(img_2_path)

        pdf = ReportPDF()
        pdf.set_auto_page_break(auto=True, margin=10)

        if img_15_path.exists():
            chart_bottom_y = add_chart_page(
                pdf=pdf,
                page_title=f"{site_name} - Last 15 Days",
                image_path=img_15_path,
                source_file_name=source_file_name,
                selected_metrics=selected_metrics,
                scale_mode=scale_mode
            )
            pdf.set_y(chart_bottom_y + 6)
            add_installation_insight_block(pdf, chart_15_df)
        else:
            add_no_data_page(
                pdf=pdf,
                page_title=f"{site_name} - Last 15 Days",
                source_file_name=source_file_name,
                selected_metrics=selected_metrics,
                scale_mode=scale_mode
            )

        if img_2_path.exists():
            chart_bottom_y = add_chart_page(
                pdf=pdf,
                page_title=f"{site_name} - Last 2 Days",
                image_path=img_2_path,
                source_file_name=source_file_name,
                selected_metrics=selected_metrics,
                scale_mode=scale_mode
            )
            pdf.set_y(chart_bottom_y + 6)
            add_installation_insight_block(pdf, chart_2_df)
        else:
            add_no_data_page(
                pdf=pdf,
                page_title=f"{site_name} - Last 2 Days",
                source_file_name=source_file_name,
                selected_metrics=selected_metrics,
                scale_mode=scale_mode
            )

        result = pdf.output(dest="S")
        return result if isinstance(result, (bytes, bytearray)) else result.encode("latin1")

    finally:
        for image_path in created_images:
            if image_path.exists():
                try:
                    image_path.unlink()
                except Exception:
                    pass


def build_optimized_excel_report(
    filtered_long,
    selected_metrics,
    scale_mode,
    source_file_name,
    site_view_states=None,
    displayed_sites=None
):
    output = BytesIO()

    if filtered_long.empty:
        return output.getvalue()

    report_timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    temp_dir = REPORT_DIR / "temp_excel_images"
    temp_dir.mkdir(parents=True, exist_ok=True)

    working_df = filtered_long.copy()
    working_df["date"] = pd.to_datetime(working_df["begin"]).dt.date
    working_df["hour"] = pd.to_datetime(working_df["begin"]).dt.hour
    working_df["month"] = pd.to_datetime(working_df["begin"]).dt.to_period("M").astype(str)
    working_df["weekday"] = pd.to_datetime(working_df["begin"]).dt.day_name()
    working_df["timewindow"] = pd.cut(
        working_df["hour"],
        bins=[-1, 5, 11, 17, 23],
        labels=["Night", "Morning", "Afternoon", "Evening"]
    )

    summary_df = (
        working_df.groupby(["site", "metric", "unit"], as_index=False)
        .agg(
            count=("value", "count"),
            avg=("value", "mean"),
            min=("value", "min"),
            max=("value", "max"),
            total=("value", "sum")
        )
        .sort_values(["site", "metric"])
        .reset_index(drop=True)
    )

    for col in ["avg", "min", "max", "total"]:
        summary_df[col] = summary_df[col].round(2)

    pivot_site_avg = (
        working_df.pivot_table(
            index="site",
            columns="metric",
            values="value",
            aggfunc="mean"
        )
        .round(2)
        .reset_index()
    )

    pivot_date_avg = (
        working_df.pivot_table(
            index="date",
            columns="metric",
            values="value",
            aggfunc="mean"
        )
        .round(2)
        .reset_index()
    )

    dashboard_info = pd.DataFrame({
        "Field": [
            "Source File",
            "Generated On",
            "Scale Mode",
            "Selected Metrics",
            "Sites",
            "Records",
            "Start Time",
            "End Time"
        ],
        "Value": [
            source_file_name,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            scale_mode,
            ", ".join(selected_metrics),
            working_df["site"].nunique(),
            len(working_df),
            str(working_df["begin"].min()),
            str(working_df["begin"].max())
        ]
    })

    chart_files = []
    chart_positions = ["A10", "A32", "A54", "A76"]
    chart_data_blocks = []

    sites_for_charts = displayed_sites if displayed_sites else sorted(working_df["site"].dropna().unique().tolist())

    for site in sites_for_charts[:4]:
        site_df = working_df[working_df["site"] == site].copy()
        if site_df.empty:
            continue

        site_state = (site_view_states or {}).get(site, {})
        current_mode = site_state.get("mode", "15 Day")

        if current_mode == "Custom":
            custom_start_date = site_state.get("custom_start_date")
            custom_end_date = site_state.get("custom_end_date")
            custom_start_time = site_state.get("custom_start_time")
            custom_end_time = site_state.get("custom_end_time")

            if all(x is not None for x in [custom_start_date, custom_end_date, custom_start_time, custom_end_time]):
                custom_start_dt = pd.Timestamp(datetime.combine(custom_start_date, custom_start_time))
                custom_end_dt = pd.Timestamp(datetime.combine(custom_end_date, custom_end_time))
                range_df = get_custom_range_df(site_df, selected_metrics, custom_start_dt, custom_end_dt)
                title_suffix = f"Custom Range | {custom_start_dt} to {custom_end_dt}"
            else:
                range_df = get_last_n_days_df(site_df, selected_metrics, 15)
                title_suffix = "Last 15 Days"
        elif current_mode == "2 Day":
            range_df = get_last_n_days_df(site_df, selected_metrics, 2)
            title_suffix = "Last 2 Days"
        else:
            range_df = get_last_n_days_df(site_df, selected_metrics, 15)
            title_suffix = "Last 15 Days"

        if range_df.empty:
            continue

        fig = build_combined_chart(
            range_df,
            site,
            scale_mode,
            chart_id=f"excel_{safe_file_name(site)}_{report_timestamp}",
            title_suffix=title_suffix
        )
        if fig is None:
            continue

        img_path = temp_dir / f"{safe_file_name(site)}_{report_timestamp}.png"
        fig.write_image(str(img_path), format="png", width=1600, height=900)
        chart_files.append(img_path)
        chart_data_blocks.append((site, img_path))

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        working_df.to_excel(writer, sheet_name="CleanData", index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        pivot_site_avg.to_excel(writer, sheet_name="PivotSiteAvg", index=False)
        pivot_date_avg.to_excel(writer, sheet_name="PivotDateAvg", index=False)
        dashboard_info.to_excel(writer, sheet_name="Dashboard", index=False, startrow=0)

    output.seek(0)
    wb = load_workbook(output)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    heat_rule = ColorScaleRule(
        start_type="min", start_color="63BE7B",
        mid_type="percentile", mid_value=50, mid_color="FFEB84",
        end_type="max", end_color="F8696B"
    )

    def style_sheet(ws, freeze="A2"):
        ws.freeze_panes = freeze
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[:300]:
                try:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 3, 28)

    for sheet_name in ["CleanData", "Summary", "PivotSiteAvg", "PivotDateAvg"]:
        ws = wb[sheet_name]
        style_sheet(ws)

    for sheet_name in ["PivotSiteAvg", "PivotDateAvg"]:
        ws = wb[sheet_name]
        if ws.max_row > 1 and ws.max_column > 1:
            ws.conditional_formatting.add(
                f"B2:{get_column_letter(ws.max_column)}{ws.max_row}",
                heat_rule
            )

    for sheet_name in ["CleanData", "Summary"]:
        ws = wb[sheet_name]
        table_ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        tab = Table(displayName=f"Tbl{sheet_name}", ref=table_ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        ws.add_table(tab)

    dashboard_ws = wb["Dashboard"]
    dashboard_ws["A1"] = "Optimized Meteorological Dashboard"
    dashboard_ws["A1"].font = Font(bold=True, size=16, color="1F1F1F")
    dashboard_ws["A2"] = "Filtered export from current dashboard selection"
    dashboard_ws["A2"].font = Font(italic=True, size=10, color="6B778C")
    dashboard_ws.column_dimensions["A"].width = 24
    dashboard_ws.column_dimensions["B"].width = 40

    for idx, (site, img_path) in enumerate(chart_data_blocks):
        if idx >= len(chart_positions):
            break
        anchor = chart_positions[idx]
        img = XLImage(str(img_path))
        img.width = 520
        img.height = 290
        dashboard_ws.add_image(img, anchor)

    final_output = BytesIO()
    wb.save(final_output)
    final_output.seek(0)

    for img_path in chart_files:
        try:
            if img_path.exists():
                img_path.unlink()
        except Exception:
            pass

    return final_output.getvalue()


# =========================================================
# MASTER TEMPLATE HELPERS
# =========================================================
def find_master_template(possible_names=None):
    if possible_names is None:
        possible_names = [
            MASTER_TEMPLATE_NAME,
            "MASTER file.xlsx",
            "MASTER_FILE.xlsx",
            "MASTER.xlsx"
        ]

    search_roots = [BASE_DIR, Path.cwd(), Path("/mnt/data")]
    search_roots = [p for p in search_roots if p.exists()]

    for root in search_roots:
        for name in possible_names:
            candidate = root / name
            if candidate.exists():
                return candidate

    for root in search_roots:
        for f in root.rglob("*.xlsx"):
            if "master" in f.name.lower():
                return f

    raise FileNotFoundError(
        f"Master workbook not found. Expected something like '{MASTER_TEMPLATE_NAME}'."
    )


def normalize_header_name(x):
    return str(x).strip().lower().replace("\n", " ").replace("_", " ")


def get_sheet_case_insensitive(wb, target_name):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == target_name.strip().lower():
            return ws
    raise KeyError(f"Sheet '{target_name}' not found. Available sheets: {wb.sheetnames}")


def build_refined_output_df(long_df):
    df = long_df.copy()

    if "begin" in df.columns:
        df["begin"] = pd.to_datetime(df["begin"], errors="coerce")
    if "end" in df.columns:
        df["end"] = pd.to_datetime(df["end"], errors="coerce")

    out = pd.DataFrame()

    if "begin" in df.columns:
        out["begin"] = df["begin"]
        out["Begin"] = df["begin"]
        out["hour"] = df["begin"].dt.hour
        out["month"] = df["begin"].dt.to_period("M").astype(str)
        out["weekday"] = df["begin"].dt.day_name()
        out["timewindow"] = pd.cut(
            df["begin"].dt.hour,
            bins=[-1, 5, 11, 17, 23],
            labels=["Night", "Morning", "Afternoon", "Evening"]
        )

    if "end" in df.columns:
        out["end"] = df["end"]
        out["End"] = df["end"]

    if "site" in df.columns:
        out["site"] = df["site"]
        out["Site"] = df["site"]

    if "metric" in df.columns:
        out["metric"] = df["metric"]
        out["Metric"] = df["metric"]

    if "unit" in df.columns:
        out["unit"] = df["unit"]
        out["Unit"] = df["unit"]

    if "parameter" in df.columns:
        out["parameter"] = df["parameter"]
        out["Parameter"] = df["parameter"]

    if "value" in df.columns:
        out["value"] = pd.to_numeric(df["value"], errors="coerce")
        out["Value"] = pd.to_numeric(df["value"], errors="coerce")

    if "begin" in df.columns:
        out["Date"] = pd.to_datetime(df["begin"], errors="coerce").dt.date
        out["Time"] = pd.to_datetime(df["begin"], errors="coerce").dt.time

    out = out.loc[:, ~out.columns.duplicated()]
    return out

def map_to_master_headers(refined_df, master_headers):
    src_cols = list(refined_df.columns)
    src_lookup = {normalize_header_name(c): c for c in src_cols}

    output_df = pd.DataFrame(index=refined_df.index)

    alias_candidates = {
        "begin": ["begin", "start", "timestamp", "datetime"],
        "end": ["end", "stop", "end time"],
        "site": ["site", "location", "station"],
        "metric": ["metric", "parameter type", "measure"],
        "unit": ["unit", "uom"],
        "parameter": ["parameter", "column", "tag"],
        "value": ["value", "reading", "measured value"],
        "date": ["date"],
        "time": ["time"]
    }

    for master_col in master_headers:
        norm_master = normalize_header_name(master_col)

        if norm_master in src_lookup:
            output_df[master_col] = refined_df[src_lookup[norm_master]]
            continue

        matched = False
        for _, aliases in alias_candidates.items():
            if norm_master in aliases or norm_master in alias_candidates:
                for alias in aliases:
                    if alias in src_lookup:
                        output_df[master_col] = refined_df[src_lookup[alias]]
                        matched = True
                        break
            if matched:
                break

        if not matched:
            output_df[master_col] = None

    return output_df


def clear_sheet_data_keep_header(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def copy_row_style(ws, source_row=2, target_row=2, max_col=None):
    if max_col is None:
        max_col = ws.max_column

    for col in range(1, max_col + 1):
        source_cell = ws.cell(row=source_row, column=col)
        target_cell = ws.cell(row=target_row, column=col)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)

        target_cell.number_format = copy(source_cell.number_format)
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)


def write_dataframe_to_master_sheet(ws, final_df):
    max_col = len(final_df.columns)
    has_template_style_row = ws.max_row >= 2

    for r_idx, row in enumerate(final_df.itertuples(index=False), start=2):
        if has_template_style_row and r_idx > 2:
            copy_row_style(ws, source_row=2, target_row=r_idx, max_col=max_col)

        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)

            if pd.isna(value):
                cell.value = None
            elif isinstance(value, pd.Timestamp):
                cell.value = value.to_pydatetime()
            else:
                cell.value = value


def refresh_excel_tables(ws):
    if not ws.tables:
        return
    table_names = list(ws.tables.keys())
    for table_name in table_names:
        tab = ws.tables[table_name]
        tab.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"


def set_workbook_calc_flags(wb):
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass


def build_master_template_report(filtered_long, source_file_name):
    if filtered_long.empty:
        return b"", pd.DataFrame(), "", []

    refined_df = build_refined_output_df(filtered_long)

    master_path = find_master_template()
    wb = load_workbook(master_path)
    ws = get_sheet_case_insensitive(wb, MASTER_SHEET_NAME)

    master_headers = [cell.value for cell in ws[1] if cell.value is not None]
    if not master_headers:
        raise ValueError(f"No headers found in row 1 of sheet '{ws.title}'.")

    final_df = map_to_master_headers(refined_df, master_headers)

    clear_sheet_data_keep_header(ws)
    write_dataframe_to_master_sheet(ws, final_df)
    refresh_excel_tables(ws)
    set_workbook_calc_flags(wb)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output.getvalue(), final_df, master_path.name, master_headers


def init_dashboard_state(current_hash, sites, metrics):
    default_metrics = [m for m in ["Wind Speed", "Wind Direction"] if m in metrics]
    if not default_metrics and metrics:
        default_metrics = metrics[:1]

    if st.session_state.get("loaded_file_hash") != current_hash:
        st.session_state["loaded_file_hash"] = current_hash
        st.session_state["applied_filters"] = {
            "scale_mode": "Normalized (0-100)",
            "selected_metrics": default_metrics,
            "displayed_sites": sites,
            "report_sites": sites
        }
        st.session_state["site_view_state"] = {}
        st.session_state["generated_site_pdfs"] = []

    st.session_state.setdefault("applied_filters", {
        "scale_mode": "Normalized (0-100)",
        "selected_metrics": default_metrics,
        "displayed_sites": sites,
        "report_sites": sites
    })
    st.session_state.setdefault("site_view_state", {})
    st.session_state.setdefault("generated_site_pdfs", {})


def ensure_site_state(site, site_min, site_max):
    if site not in st.session_state["site_view_state"]:
        st.session_state["site_view_state"][site] = {
            "mode": "15 Day",
            "custom_start_date": site_min.date() if pd.notna(site_min) else datetime.now().date(),
            "custom_end_date": site_max.date() if pd.notna(site_max) else datetime.now().date(),
            "custom_start_time": time(0, 0),
            "custom_end_time": time(23, 59)
        }


# =========================================================
# UI
# =========================================================
st.markdown("""
<div class="main-shell">
    <div class="hero-panel">
        <div class="hero-title">Meteorological Operations Suite</div>
        <div class="hero-subtitle">
            Compact operational monitoring dashboard for multi-site meteorological data.
            Focused on clean charting, fast filtering, and minimal on-screen clutter.
        </div>
        <span class="pill">Compact UI</span>
        <span class="pill">Instant 15/2 Days</span>
        <span class="pill">Custom Range View</span>
        <span class="pill">Per-Site PDF Reports</span>
        <span class="pill">Master Template Export</span>
    </div>
</div>
""", unsafe_allow_html=True)

st.sidebar.markdown("### Data Input")
uploaded_file = st.sidebar.file_uploader("Upload CSV / Excel", type=["csv", "xlsx", "xls"])
show_summary_tables = st.sidebar.toggle("Show metric summaries", value=False)
show_raw = st.sidebar.toggle("Show processed raw data", value=False)
show_upload_history = st.sidebar.toggle("Show upload history", value=False)

if uploaded_file is None:
    st.info("Upload your latest meteorological file to begin.")
    if show_upload_history:
        log_df = load_upload_log()
        if not log_df.empty:
            st.dataframe(
                log_df.sort_values("upload_timestamp", ascending=False).head(10),
                use_container_width=True,
                hide_index=True
            )
    st.stop()

try:
    file_bytes = uploaded_file.getvalue()
    with st.spinner("Processing uploaded file..."):
        wide_df, long_df, begin_col, end_col = parse_uploaded_content(file_bytes, uploaded_file.name)
        long_df = prepare_long_df(long_df)
except Exception as e:
    st.error(f"File could not be parsed. Error: {e}")
    st.stop()

current_hash = get_file_hash(file_bytes)
if st.session_state.get("last_saved_hash") != current_hash:
    save_upload_bundle(uploaded_file, file_bytes, wide_df, long_df)
    st.session_state["last_saved_hash"] = current_hash
    st.success(f"Uploaded and saved: {uploaded_file.name}")

sites = sorted(long_df["site"].dropna().unique().tolist())
metrics = sorted(long_df["metric"].dropna().unique().tolist())

init_dashboard_state(current_hash, sites, metrics)
applied = st.session_state["applied_filters"]

st.sidebar.markdown("### Analysis Controls")
with st.sidebar.form("dashboard_filter_form"):
    form_scale_mode = st.selectbox(
        "Combined graph scale",
        ["Normalized (0-100)", "Actual Values"],
        index=["Normalized (0-100)", "Actual Values"].index(applied["scale_mode"])
        if applied["scale_mode"] in ["Normalized (0-100)", "Actual Values"] else 0
    )
    form_selected_metrics = st.multiselect(
        "Parameters",
        options=metrics,
        default=[m for m in applied["selected_metrics"] if m in metrics]
    )
    form_displayed_sites = st.multiselect(
        "Sites",
        options=sites,
        default=[s for s in applied["displayed_sites"] if s in sites]
    )
    form_report_sites = st.multiselect(
        "Report sites",
        options=sites,
        default=[s for s in applied["report_sites"] if s in sites]
    )

    view_dashboard = st.form_submit_button("View Dashboard", use_container_width=True)

if view_dashboard:
    st.session_state["applied_filters"] = {
        "scale_mode": form_scale_mode,
        "selected_metrics": form_selected_metrics,
        "displayed_sites": form_displayed_sites,
        "report_sites": form_report_sites
    }

applied = st.session_state["applied_filters"]
scale_mode = applied["scale_mode"]
selected_metrics = applied["selected_metrics"]
displayed_sites = applied["displayed_sites"]
selected_sites_for_reports = applied["report_sites"]

filtered_long = long_df[
    long_df["site"].isin(selected_sites_for_reports) &
    long_df["metric"].isin(selected_metrics)
].copy()

summary_tables = {}
if show_summary_tables and selected_metrics:
    st.markdown('<div class="section-label">Metric Summaries</div>', unsafe_allow_html=True)
    for metric in selected_metrics:
        metric_df = filtered_long[filtered_long["metric"] == metric].copy()
        if metric_df.empty:
            continue

        unit_mode = metric_df["unit"].dropna().astype(str).mode()
        unit = unit_mode.iloc[0] if len(unit_mode) > 0 else ""
        summary_df = make_metric_summary(metric_df)
        summary_tables[metric] = summary_df
        insights = build_insights(metric_df, metric, unit)

        with st.expander(f"{metric} summary", expanded=False):
            for item in insights:
                st.write(f"- {item}")
            st.dataframe(summary_df, use_container_width=True, hide_index=True)

st.markdown("""
<div class="subtle-note">
Apply filters from the sidebar, then use <b>View Dashboard</b>.
Each site card supports 15 Day, 2 Day, and Custom views.
The PDF generator creates one PDF per selected report site and includes exactly two charts per PDF: 15 Day and 2 Day.
</div>
""", unsafe_allow_html=True)

st.markdown('<div class="section-label">Site Analysis</div>', unsafe_allow_html=True)

if not selected_metrics:
    st.warning("Please select at least one parameter and click View Dashboard.")
else:
    for site in displayed_sites:
        site_df = long_df[long_df["site"] == site].copy()
        if site_df.empty:
            continue

        site_min = site_df["begin"].min()
        site_max = site_df["begin"].max()
        ensure_site_state(site, site_min, site_max)

        site_state = st.session_state["site_view_state"][site]
        safe_site = safe_file_name(site)

        metrics_present = sorted(site_df[site_df["metric"].isin(selected_metrics)]["metric"].dropna().unique().tolist())
        chips_html = "".join([f'<span class="metric-chip">{m}</span>' for m in metrics_present])

        st.markdown('<div class="site-shell">', unsafe_allow_html=True)
        st.markdown(
            f"""
            <div class="site-title">{site}</div>
            <div class="site-subtitle">
                Available range: {str(site_min)[:16] if pd.notna(site_min) else "-"} to {str(site_max)[:16] if pd.notna(site_max) else "-"}
            </div>
            <div style="margin-top:4px;">{chips_html}</div>
            """,
            unsafe_allow_html=True
        )

        current_mode = st.radio(
            "View range",
            options=["15 Day", "2 Day", "Custom"],
            index=["15 Day", "2 Day", "Custom"].index(site_state["mode"])
            if site_state["mode"] in ["15 Day", "2 Day", "Custom"] else 0,
            horizontal=True,
            key=f"{safe_site}_mode_radio",
            label_visibility="collapsed"
        )

        site_state["mode"] = current_mode
        st.markdown(
            f'<div style="margin:6px 0 10px 0;"><span class="range-badge">Current View: {site_state["mode"]}</span></div>',
            unsafe_allow_html=True
        )

        if current_mode == "Custom":
            st.markdown(
                '<div class="helper-note">Choose dates and time, then click <b>View Custom Range</b>.</div>',
                unsafe_allow_html=True
            )

            with st.form(f"custom_form_{safe_site}"):
                c1, c2, c3, c4 = st.columns(4)

                with c1:
                    custom_start_date = st.date_input(
                        "Start date",
                        value=site_state["custom_start_date"],
                        min_value=site_min.date() if pd.notna(site_min) else None,
                        max_value=site_max.date() if pd.notna(site_max) else None,
                        key=f"{safe_site}_custom_start_date_widget"
                    )
                with c2:
                    custom_end_date = st.date_input(
                        "End date",
                        value=site_state["custom_end_date"],
                        min_value=site_min.date() if pd.notna(site_min) else None,
                        max_value=site_max.date() if pd.notna(site_max) else None,
                        key=f"{safe_site}_custom_end_date_widget"
                    )
                with c3:
                    custom_start_time = st.time_input(
                        "Start time",
                        value=site_state["custom_start_time"],
                        key=f"{safe_site}_custom_start_time_widget"
                    )
                with c4:
                    custom_end_time = st.time_input(
                        "End time",
                        value=site_state["custom_end_time"],
                        key=f"{safe_site}_custom_end_time_widget"
                    )

                view_custom = st.form_submit_button("View Custom Range", use_container_width=True)

            if view_custom:
                site_state["custom_start_date"] = custom_start_date
                site_state["custom_end_date"] = custom_end_date
                site_state["custom_start_time"] = custom_start_time
                site_state["custom_end_time"] = custom_end_time

            custom_start_dt = pd.Timestamp(datetime.combine(site_state["custom_start_date"], site_state["custom_start_time"]))
            custom_end_dt = pd.Timestamp(datetime.combine(site_state["custom_end_date"], site_state["custom_end_time"]))

            if custom_end_dt < custom_start_dt:
                st.warning("Custom range end must be after start.")
                range_df = pd.DataFrame()
                title_suffix = "Custom Range"
            else:
                range_df = get_custom_range_df(site_df, selected_metrics, custom_start_dt, custom_end_dt)
                title_suffix = f"Custom Range | {custom_start_dt} to {custom_end_dt}"

        elif current_mode == "2 Day":
            range_df = get_last_n_days_df(site_df, selected_metrics, 2)
            title_suffix = "Last 2 Days"

        else:
            range_df = get_last_n_days_df(site_df, selected_metrics, 15)
            title_suffix = "Last 15 Days"

        fig = build_combined_chart(
            range_df,
            site,
            scale_mode,
            chart_id=f"{safe_site}_{current_mode}_{scale_mode}",
            title_suffix=title_suffix
        )

        if fig is not None:
            st.plotly_chart(
                fig,
                use_container_width=True,
                config={"displaylogo": False, "responsive": True}
            )
        else:
            st.info("No data available for the current site view.")

        latest_df = build_latest_table(range_df)
        summary_df = build_window_summary(range_df)

        tab1, tab2 = st.tabs(["Latest values", "Window summary"])
        with tab1:
            if latest_df.empty:
                st.caption("No latest values available for this view.")
            else:
                st.dataframe(latest_df, use_container_width=True, hide_index=True)

        with tab2:
            if summary_df.empty:
                st.caption("No summary available for this view.")
            else:
                st.dataframe(summary_df, use_container_width=True, hide_index=True)

        st.markdown('</div>', unsafe_allow_html=True)

st.markdown('<div class="section-label">Generate Site PDFs</div>', unsafe_allow_html=True)

if st.button("Generate Site PDF Reports"):
    if not selected_metrics:
        st.warning("Please select at least one parameter before generating PDFs.")
    elif not selected_sites_for_reports:
        st.warning("Please select at least one report site before generating PDFs.")
    else:
        generated_reports = []
        report_timestamp = datetime.now().strftime("%Y%m%d%H%M%S")

        with st.spinner("Generating site PDF reports..."):
            for site in selected_sites_for_reports:
                site_df = long_df[
                    (long_df["site"] == site) &
                    (long_df["metric"].isin(selected_metrics))
                ].copy()

                if site_df.empty:
                    continue

                try:
                    pdf_bytes = build_site_pdf_report_bytes(
                        site_name=site,
                        site_df=site_df,
                        selected_metrics=selected_metrics,
                        scale_mode=scale_mode,
                        source_file_name=uploaded_file.name,
                        report_timestamp=report_timestamp
                    )
                    pdf_name = f"{safe_file_name(site)}_{report_timestamp}.pdf"
                    pdf_path = REPORT_DIR / pdf_name

                    with open(pdf_path, "wb") as f:
                        f.write(pdf_bytes)

                    generated_reports.append({
                        "site": site,
                        "file_name": pdf_name,
                        "path": str(pdf_path)
                    })

                except Exception as e:
                    st.error(f"Could not generate PDF for {site}: {e}")

        st.session_state["generated_site_pdfs"] = generated_reports

        if generated_reports:
            st.success(f"Generated {len(generated_reports)} site PDF reports in REPORT_DIR.")
        else:
            st.warning("No site PDFs were generated.")

st.markdown("### Generate Optimized Excel")

if st.button("Generate Optimized Excel"):
    if filtered_long.empty:
        st.warning("No filtered data available to export.")
    elif not selected_metrics:
        st.warning("Please select at least one parameter before generating the Excel report.")
    else:
        with st.spinner("Generating optimized Excel report..."):
            excel_bytes = build_optimized_excel_report(
                filtered_long=filtered_long,
                selected_metrics=selected_metrics,
                scale_mode=scale_mode,
                source_file_name=uploaded_file.name,
                site_view_states=st.session_state.get("site_view_state"),
                displayed_sites=displayed_sites
            )

        excel_name = f"Optimized_Meteo_Report_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        st.download_button(
            label="Download Optimized Excel",
            data=excel_bytes,
            file_name=excel_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

st.markdown("### Generate Master Template Excel")

if st.button("Generate Master Template Excel"):
    if filtered_long.empty:
        st.warning("No filtered data available to export into the master template.")
    else:
        try:
            with st.spinner("Populating master template workbook..."):
                master_bytes, master_preview_df, master_filename, master_headers = build_master_template_report(
                    filtered_long=filtered_long,
                    source_file_name=uploaded_file.name
                )

            st.success(f"Master workbook '{master_filename}' populated successfully into sheet '{MASTER_SHEET_NAME}'.")

            with st.expander("Preview mapped master data", expanded=False):
                st.write(master_headers)
                st.dataframe(master_preview_df, use_container_width=True, hide_index=True)

            master_output_name = f"Updated_MASTER_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
            st.download_button(
                label="Download Updated Master Workbook",
                data=master_bytes,
                file_name=master_output_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

        except Exception as e:
            st.error(f"Master template export failed: {e}")

generated_site_pdfs = st.session_state.get("generated_site_pdfs", [])

if generated_site_pdfs:
    st.markdown('<div class="section-label">PDF Downloads</div>', unsafe_allow_html=True)

    for report in generated_site_pdfs:
        c1, c2 = st.columns([2.4, 1])

        with c1:
            st.write(report["site"])
            st.caption(report["path"])

        with c2:
            pdf_path = report["path"]
            with open(pdf_path, "rb") as pdf_file:
                pdf_data = pdf_file.read()

            st.download_button(
                label=f"Download {report['site']} PDF",
                data=pdf_data,
                file_name=report["file_name"],
                mime="application/pdf",
                key=f"download_{safe_file_name(report['site'])}_{report['file_name']}",
                use_container_width=True
            )

    st.markdown("### Email Reports")
    email_recipients_input = st.text_area("Recipient emails comma separated", value="")
    email_subject = st.text_input("Email subject", value="Meteorological Site Reports")
    email_body = st.text_area(
        "Email body",
        value="Please find attached the generated meteorological site PDF reports."
    )

    if st.button("Send Email With All Reports"):
        if not generated_site_pdfs:
            st.warning("No generated reports available to email.")
        else:
            to_emails = [e.strip() for e in email_recipients_input.split(",") if e.strip()]
            if not to_emails:
                st.warning("Please enter at least one recipient email.")
            else:
                try:
                    send_reports_email_gmail(
                        to_emails=to_emails,
                        subject=email_subject,
                        body=email_body,
                        reports=generated_site_pdfs
                    )
                    st.success("Email sent successfully with all attached reports.")
                except Exception as e:
                    st.error(f"Email failed: {e}")

if show_raw:
    st.markdown('<div class="section-label">Processed Raw Data</div>', unsafe_allow_html=True)
    st.dataframe(filtered_long, use_container_width=True, hide_index=True)

if show_upload_history:
    log_df = load_upload_log()
    if not log_df.empty:
        st.markdown('<div class="section-label">Upload History</div>', unsafe_allow_html=True)
        st.dataframe(
            log_df.sort_values("upload_timestamp", ascending=False).head(20),
            use_container_width=True,
            hide_index=True
        )
